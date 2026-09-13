from __future__ import annotations

import io

import pandas as pd
from fastapi.responses import StreamingResponse

import main


# Replace the legacy tracker export with an export that backfills deterministic
# AI assessment fields when a candidate has both a stored resume and assigned JD.
main.app.router.routes = [
    route
    for route in main.app.router.routes
    if not (
        getattr(route, "path", None) == "/api/export/tracker"
        and "GET" in (getattr(route, "methods", set()) or set())
    )
]


def _as_list(value):
    if isinstance(value, list):
        return [str(v) for v in value if str(v).strip()]
    if value is None or value == "":
        return []
    return [str(value)]


def _ai_feedback(a: dict) -> str:
    evidence = _as_list(a.get("evidence"))
    risks = _as_list(a.get("risks"))
    parts = []
    if evidence:
        parts.append("Evidence: " + " ".join(evidence))
    if risks:
        parts.append("Verify: " + " ".join(risks))
    return " | ".join(parts)


def _assessment_for_candidate(candidate: dict) -> dict:
    stored = main._decode_ai_details(candidate.get("ai_details")) or {}
    useful_keys = (
        "score",
        "rating",
        "skill_coverage",
        "experience_fit",
        "matched_skills",
        "missing_skills",
        "risks",
        "evidence",
        "screening_questions",
    )
    if any(stored.get(key) not in (None, "", [], {}) for key in useful_keys):
        return stored

    resume = candidate.get("resume_text") or ""
    jd = candidate.get("job_jd") or ""
    if not resume.strip() or not jd.strip():
        return stored

    sim = main.semantic_scores(jd, [resume])[0]
    return main.score_resume(jd, resume, sim)


@main.app.get("/api/export/tracker")
def export_candidate_tracker():
    con = main.db()
    rows = con.execute(
        """SELECT c.*,j.title job_title,j.department job_department,j.jd job_jd
           FROM candidates c
           LEFT JOIN jobs j ON j.id=c.job_id
           ORDER BY c.id ASC"""
    ).fetchall()
    con.close()

    export_rows = []
    for i, row in enumerate(rows, 1):
        c = dict(row)
        p = main._decode_ai_details(c.get("profile_details")) or {}
        a = _assessment_for_candidate(c)
        created = (c.get("created_at") or "")[:10]
        matched = _as_list(a.get("matched_skills"))
        missing = _as_list(a.get("missing_skills"))
        risks = _as_list(a.get("risks"))
        questions = _as_list(a.get("screening_questions"))

        export_rows.append({
            "S No": i,
            "Date": created,
            "Client": p.get("client", ""),
            "Position Worked": c.get("job_title") or p.get("role_name", ""),
            "Tech Stack": c.get("skills") or "",
            "Recruter": p.get("recruiter_name", ""),
            "Resource Name": c.get("name") or "",
            "Number": c.get("phone") or "",
            "Mail-id": c.get("email") or "",
            "Current Company": p.get("current_organization", ""),
            "Current Designation": p.get("current_designation", ""),
            "Tot Exp (In Years)": c.get("experience") if c.get("experience") is not None else p.get("total_experience", ""),
            "Rel Exp (In Years)": p.get("relevant_experience", ""),
            "Current company Exp (In Years)": p.get("current_company_experience", ""),
            "CTC (In LPA)": c.get("current_ctc") or p.get("current_ctc", ""),
            "ECTC (In LPA)": c.get("expected_ctc") or p.get("expected_ctc", ""),
            "Offer In hand (if any)": p.get("holding_offers", ""),
            "Last Appraisal & Month": p.get("last_appraisal", ""),
            "Notice Period": c.get("notice_period") or p.get("notice_period", ""),
            "Native Location": p.get("native_location", ""),
            "Current Location": p.get("current_location", ""),
            "Preferred Location": p.get("preferred_location", ""),
            "UG": p.get("ug", ""),
            "Highest Qualification": p.get("highest_qualification", ""),
            "LinkedIn ID": p.get("linkedin_id", ""),
            "Profile Submission Date": created,
            "Screening Status": p.get("screening_status", ""),
            "Interview Level": p.get("interview_level", ""),
            "Status": c.get("stage") or "",
            "Offer": p.get("offer", ""),
            "D.O.J": p.get("tentative_doj", ""),
            "Remarks": p.get("remarks", ""),
            "AI Score": c.get("ai_score") if c.get("ai_score") is not None else a.get("score", ""),
            "AI Rating": c.get("rating") or a.get("rating", ""),
            "AI Feedback": _ai_feedback(a),
            "Skill Match %": a.get("skill_coverage", ""),
            "Experience Fit %": a.get("experience_fit", ""),
            "Matched Skills": ", ".join(matched),
            "Missing Skills": ", ".join(missing),
            "Risks / Points to Verify": " | ".join(risks),
            "Suggested Screening Questions": " | ".join(questions),
        })

    out = io.BytesIO()
    df = pd.DataFrame(export_rows)
    if df.empty:
        df = pd.DataFrame(columns=[
            "S No","Date","Client","Position Worked","Tech Stack","Recruter","Resource Name","Number","Mail-id",
            "Current Company","Current Designation","Tot Exp (In Years)","Rel Exp (In Years)","Current company Exp (In Years)",
            "CTC (In LPA)","ECTC (In LPA)","Offer In hand (if any)","Last Appraisal & Month","Notice Period","Native Location","Current Location",
            "Preferred Location","UG","Highest Qualification","LinkedIn ID","Profile Submission Date","Screening Status",
            "Interview Level","Status","Offer","D.O.J","Remarks","AI Score","AI Rating","AI Feedback","Skill Match %","Experience Fit %",
            "Matched Skills","Missing Skills","Risks / Points to Verify","Suggested Screening Questions"
        ])

    with pd.ExcelWriter(out, engine="openpyxl") as writer:
        df.to_excel(writer, index=False, sheet_name="Candidate Tracker")
        ws = writer.book["Candidate Tracker"]
        from openpyxl.styles import PatternFill, Font, Alignment

        header_fill = PatternFill("solid", fgColor="000000")
        ai_fill = PatternFill("solid", fgColor="C00000")
        white_font = Font(color="FFFFFF", bold=True)
        for cell in ws[1]:
            cell.fill = ai_fill if cell.column >= 33 else header_fill
            cell.font = white_font
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        for col in ws.columns:
            letter = col[0].column_letter
            max_len = max(len(str(cell.value or "")) for cell in col[:100])
            ws.column_dimensions[letter].width = min(max(max_len + 2, 12), 38)
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=True)

    out.seek(0)
    return StreamingResponse(
        out,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=ShortlistAI_Candidate_Tracker.xlsx"},
    )


app = main.app
