from pathlib import Path

main = Path('main.py')
s = main.read_text(encoding='utf-8')

# 1) Make education parsing work with both abbreviations and full degree names.
start = s.index('def _detect_education(')
end = s.index('\ndef extract_profile_details(', start)
new_education = r'''def _detect_education(text: str) -> Dict[str,str]:
    patterns = [
        (r"(?i)\b(?:ph\.?d|doctor(?:ate)?\s+of\s+philosophy)\b", "PHD", "PG"),
        (r"(?i)\b(?:m\.?tech|master\s+of\s+technology)\b", "MTECH", "PG"),
        (r"(?i)\b(?:m\.?e\.?|master\s+of\s+engineering)\b", "ME", "PG"),
        (r"(?i)\b(?:mca|master\s+of\s+computer\s+applications?)\b", "MCA", "PG"),
        (r"(?i)\b(?:mba|master\s+of\s+business\s+administration)\b", "MBA", "PG"),
        (r"(?i)\b(?:m\.?sc|master\s+of\s+science)\b", "MSC", "PG"),
        (r"(?i)\b(?:b\.?tech|bachelor\s+of\s+technology)\b", "BTECH", "UG"),
        (r"(?i)\b(?:b\.?e\.?|bachelor\s+of\s+engineering)\b", "BE", "UG"),
        (r"(?i)\b(?:bca|bachelor\s+of\s+computer\s+applications?)\b", "BCA", "UG"),
        (r"(?i)\b(?:b\.?sc|bachelor\s+of\s+science)\b", "BSC", "UG"),
        (r"(?i)\b(?:b\.?com|bachelor\s+of\s+commerce)\b", "BCOM", "UG"),
        (r"(?i)\b(?:b\.?a\.?|bachelor\s+of\s+arts)\b", "BA", "UG"),
    ]
    found = []
    for pattern, code, level in patterns:
        if re.search(pattern, text):
            found.append((code, level))
    ug = next((code for code, level in found if level == "UG"), "")
    pg = next((code for code, level in found if level == "PG"), "")
    return {"ug": ug, "highest_qualification": pg or ug}


def _detect_location(text: str) -> str:
    labelled = _label_value(text, [r"current\s*location", r"present\s*location", r"location"], 70)
    if labelled:
        return labelled
    cities = [
        "Chennai","Bengaluru","Bangalore","Hyderabad","Pune","Mumbai","Delhi","New Delhi",
        "Noida","Gurugram","Gurgaon","Kolkata","Coimbatore","Kochi","Ahmedabad","Indore",
        "Bhopal","Jaipur","Trivandrum","Thiruvananthapuram","Mysuru","Mysore"
    ]
    top = "\n".join([x.strip() for x in text.splitlines()[:30] if x.strip()])
    for city in cities:
        if re.search(r"(?i)(?<![A-Za-z])" + re.escape(city) + r"(?![A-Za-z])", top):
            return "Bengaluru" if city.lower() == "bangalore" else "Gurugram" if city.lower() == "gurgaon" else city
    return ""


def _latest_experience(text: str) -> Dict[str,str]:
    lines = [re.sub(r"\s+", " ", x).strip(" \t•") for x in text.splitlines()]
    date_re = re.compile(r"(?i)\b(?:0?[1-9]|1[0-2])[/\-.](?:19|20)\d{2}\s*(?:-|–|to)\s*(?:(?:0?[1-9]|1[0-2])[/\-.](?:19|20)\d{2}|present|current)\b")
    for i, line in enumerate(lines):
        if not line or not date_re.search(line):
            continue
        title = date_re.sub("", line).strip(" ,-–|")
        title = re.sub(r"\s+,\s*$", "", title).strip()
        company_line = ""
        for j in range(i + 1, min(i + 5, len(lines))):
            cand = lines[j].strip()
            if not cand or cand.lower() in {"experience", "work experience", "professional experience"}:
                continue
            if cand.startswith(("•", "-")):
                continue
            company_line = cand
            break
        company, location = company_line, ""
        if company_line and "," in company_line:
            parts = [p.strip() for p in company_line.split(",") if p.strip()]
            if parts:
                company = parts[0]
                if len(parts) > 1:
                    location = parts[-1]
        return {"current_designation": title, "current_organization": company, "current_location": location}
    return {"current_designation":"", "current_organization":"", "current_location":""}
'''
s = s[:start] + new_education + s[end:]

# 2) Improve profile extraction using latest employment and top-of-resume location clues.
start = s.index('def extract_profile_details(')
end = s.index('\ndef semantic_scores(', start)
new_profile = r'''def extract_profile_details(text: str, filename: str) -> Dict[str,Any]:
    edu = _detect_education(text)
    latest = _latest_experience(text)
    linkedin = ""
    m = re.search(r"(?i)(?:https?://)?(?:www\.)?linkedin\.com/in/[A-Za-z0-9_\-%/]+", text)
    if m:
        linkedin = m.group(0)
        if not linkedin.lower().startswith("http"):
            linkedin = "https://" + linkedin
    notice = _label_value(text, [r"notice\s*period", r"np"], 45)
    if not notice and re.search(r"(?i)\bimmediate(?:ly)?\s+(?:available|joiner|joining)\b|\bimmediate joiner\b", text):
        notice = "Immediate"
    current_org = _label_value(text, [r"current\s*(?:organization|organisation|company)", r"present\s*(?:organization|organisation|company)"], 90) or latest["current_organization"]
    current_role = _label_value(text, [r"current\s*(?:designation|role|title)", r"designation", r"job\s*title"], 90) or latest["current_designation"]
    current_loc = _label_value(text, [r"current\s*location", r"present\s*location"], 60) or latest["current_location"] or _detect_location(text)
    summary = ""
    sm = re.search(r"(?is)\b(?:summary|professional summary|profile summary)\b\s*[:\-]?\s*(.{40,800}?)(?=\n\s*(?:skills?|experience|employment|education|projects?|certifications?)\b)", text)
    if sm:
        summary = re.sub(r"\s+", " ", sm.group(1)).strip()[:700]
    details = {
        "candidate_full_name": detect_name(text, filename),
        "contact_no": detect_phone(text),
        "email_id": detect_email(text),
        "total_experience": parse_candidate_years(text),
        "relevant_experience": _label_value(text, [r"relevant\s*(?:experience|exp)"], 40),
        "current_organization": current_org,
        "current_designation": current_role,
        "current_company_experience": _label_value(text, [r"current\s*company\s*(?:experience|exp)"], 40),
        "current_ctc": _money_value(text, [r"current\s*ctc", r"present\s*ctc"]),
        "expected_ctc": _money_value(text, [r"expected\s*ctc", r"expecting\s*ctc"]),
        "holding_offers": _label_value(text, [r"holding\s*offers?", r"offers?\s*in\s*hand", r"offer\s*in\s*hand"], 80),
        "notice_period": notice,
        "lwd": _label_value(text, [r"lwd", r"last\s*working\s*day"], 45),
        "tentative_doj": _label_value(text, [r"tentative\s*doj", r"date\s*of\s*joining", r"doj"], 45),
        "native_location": _label_value(text, [r"native\s*location", r"native\s*place"], 60),
        "current_location": current_loc,
        "preferred_location": _label_value(text, [r"preferred\s*location", r"preferred\s*locations"], 100),
        "linkedin_id": linkedin,
        "profile_link": linkedin,
        "ug": edu["ug"],
        "highest_qualification": edu["highest_qualification"],
        "skills": ", ".join(find_skills(text)),
        "profile_summary": summary,
    }
    return details
'''
s = s[:start] + new_profile + s[end:]

# 3) Add duplicate candidate detection.
if 'def _candidate_duplicate(' not in s:
    anchor = '''def _candidate_dict(row: sqlite3.Row) -> Dict[str,Any]:\n    d = dict(row)\n    d["ai_evaluation"] = _decode_ai_details(d.pop("ai_details", None))\n    d["profile_details"] = _decode_ai_details(d.get("profile_details")) or {}\n    return d\n'''
    addition = r'''

def _candidate_duplicate(con: sqlite3.Connection, email: str = "", phone: str = ""):
    email = (email or "").strip().lower()
    phone_digits = re.sub(r"\D", "", phone or "")[-10:]
    if email:
        row = con.execute("SELECT id,name,email,phone FROM candidates WHERE lower(trim(email))=? LIMIT 1", (email,)).fetchone()
        if row:
            return row
    if phone_digits:
        rows = con.execute("SELECT id,name,email,phone FROM candidates WHERE phone IS NOT NULL AND phone<>''").fetchall()
        for row in rows:
            if re.sub(r"\D", "", row["phone"] or "")[-10:] == phone_digits:
                return row
    return None
'''
    assert anchor in s
    s = s.replace(anchor, anchor + addition, 1)

needle = '''    con = db()\n    cur = con.cursor()\n    cur.execute(\n        """INSERT INTO candidates(name,email,phone,experience,skills,resume_text,resume_filename,source,notice_period,current_ctc,\n'''
if 'Possible duplicate candidate' not in s:
    replacement = '''    con = db()\n    duplicate = _candidate_duplicate(con, x.email, x.phone)\n    if duplicate:\n        name = duplicate["name"] or "Existing candidate"\n        con.close()\n        raise HTTPException(409, f"Possible duplicate candidate: {name} already exists in the ATS.")\n    cur = con.cursor()\n    cur.execute(\n        """INSERT INTO candidates(name,email,phone,experience,skills,resume_text,resume_filename,source,notice_period,current_ctc,\n'''
    assert needle in s
    s = s.replace(needle, replacement, 1)

# 4) Skip duplicate saves from AI Match and report them.
if 'duplicates_skipped = []' not in s:
    s = s.replace('''    results = []\n    now = datetime.utcnow().isoformat()\n    con = db() if save_to_ats else None\n''', '''    results = []\n    duplicates_skipped = []\n    now = datetime.utcnow().isoformat()\n    con = db() if save_to_ats else None\n''', 1)
    old = '''        if con:\n            skills = ", ".join(find_skills(txt))\n            con.execute(\n                """INSERT INTO candidates(name,email,phone,experience,skills,resume_text,resume_filename,\n                   source,job_id,stage,ai_score,rating,ai_details,profile_details,created_at,updated_at)\n                   VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)""",\n                (s["candidate"], detect_email(txt), detect_phone(txt), s["candidate_years"], skills,\n                 txt, name, "Resume upload", job_id, "Sourced", s["score"], s["rating"],\n                 json.dumps(s), json.dumps(extract_profile_details(txt, name)), now, now)\n            )\n'''
    new = '''        if con:\n            email = detect_email(txt)\n            phone = detect_phone(txt)\n            duplicate = _candidate_duplicate(con, email, phone)\n            if duplicate:\n                duplicates_skipped.append({"candidate": s["candidate"], "existing_id": duplicate["id"], "existing_name": duplicate["name"]})\n                continue\n            skills = ", ".join(find_skills(txt))\n            con.execute(\n                """INSERT INTO candidates(name,email,phone,experience,skills,resume_text,resume_filename,\n                   source,job_id,stage,ai_score,rating,ai_details,profile_details,created_at,updated_at)\n                   VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)""",\n                (s["candidate"], email, phone, s["candidate_years"], skills,\n                 txt, name, "Resume upload", job_id, "Sourced", s["score"], s["rating"],\n                 json.dumps(s), json.dumps(extract_profile_details(txt, name)), now, now)\n            )\n'''
    assert old in s
    s = s.replace(old, new, 1)
    s = s.replace('''            "weak": sum(r["rating"] == "Weak" for r in results),\n        },\n        "results": results,\n        "errors": errors,\n''', '''            "weak": sum(r["rating"] == "Weak" for r in results),\n            "duplicates_skipped": len(duplicates_skipped),\n        },\n        "results": results,\n        "errors": errors,\n        "duplicates_skipped": duplicates_skipped,\n''', 1)

# 5) Add a full ATS tracker Excel export using the user's tracker structure, with AI fields appended.
if '@app.get("/api/export/tracker")' not in s:
    endpoint = r'''

@app.get("/api/export/tracker")
def export_candidate_tracker():
    con = db()
    rows = con.execute("""SELECT c.*,j.title job_title,j.department job_department
                          FROM candidates c LEFT JOIN jobs j ON j.id=c.job_id
                          ORDER BY c.id ASC""").fetchall()
    con.close()
    export_rows = []
    for i, row in enumerate(rows, 1):
        c = dict(row)
        p = _decode_ai_details(c.get("profile_details")) or {}
        a = _decode_ai_details(c.get("ai_details")) or {}
        created = (c.get("created_at") or "")[:10]
        export_rows.append({
            "S No": i,
            "Date": created,
            "Client": p.get("client", ""),
            "Position Worked": c.get("job_title") or p.get("role_name", ""),
            "Tech Stack": c.get("skills") or "",
            "Recruiter": p.get("recruiter_name", ""),
            "Resource Name": c.get("name") or "",
            "Contact Number": c.get("phone") or "",
            "Email": c.get("email") or "",
            "Current Company": p.get("current_organization", ""),
            "Current Designation": p.get("current_designation", ""),
            "Total Experience": c.get("experience") if c.get("experience") is not None else p.get("total_experience", ""),
            "Relevant Experience": p.get("relevant_experience", ""),
            "Current Company Experience": p.get("current_company_experience", ""),
            "Current CTC": c.get("current_ctc") or p.get("current_ctc", ""),
            "Expected CTC": c.get("expected_ctc") or p.get("expected_ctc", ""),
            "Offer in Hand": p.get("holding_offers", ""),
            "Last Appraisal": p.get("last_appraisal", ""),
            "Notice Period": c.get("notice_period") or p.get("notice_period", ""),
            "Native Location": p.get("native_location", ""),
            "Current Location": p.get("current_location", ""),
            "Preferred Location": p.get("preferred_location", ""),
            "UG": p.get("ug", ""),
            "Highest Qualification": p.get("highest_qualification", ""),
            "LinkedIn ID": p.get("linkedin_id", ""),
            "Profile Submission Date": created,
            "Screening Status": p.get("screening_status", ""),
            "Interview Level": p.get("interview_level", ""),
            "Status": c.get("stage") or "",
            "Offer": p.get("offer", ""),
            "DOJ": p.get("tentative_doj", ""),
            "Remarks": p.get("remarks", ""),
            "AI Score": c.get("ai_score") if c.get("ai_score") is not None else a.get("score", ""),
            "AI Rating": c.get("rating") or a.get("rating", ""),
            "Skill Match %": a.get("skill_coverage", ""),
            "Experience Fit %": a.get("experience_fit", ""),
            "Matched Skills": ", ".join(a.get("matched_skills", [])),
            "Missing Skills": ", ".join(a.get("missing_skills", [])),
            "Risks / Points to Verify": " | ".join(a.get("risks", [])),
            "Suggested Screening Questions": " | ".join(a.get("screening_questions", [])),
        })
    out = io.BytesIO()
    df = pd.DataFrame(export_rows)
    if df.empty:
        df = pd.DataFrame(columns=[
            "S No","Date","Client","Position Worked","Tech Stack","Recruiter","Resource Name","Contact Number","Email",
            "Current Company","Current Designation","Total Experience","Relevant Experience","Current Company Experience",
            "Current CTC","Expected CTC","Offer in Hand","Last Appraisal","Notice Period","Native Location","Current Location",
            "Preferred Location","UG","Highest Qualification","LinkedIn ID","Profile Submission Date","Screening Status",
            "Interview Level","Status","Offer","DOJ","Remarks","AI Score","AI Rating","Skill Match %","Experience Fit %",
            "Matched Skills","Missing Skills","Risks / Points to Verify","Suggested Screening Questions"
        ])
    with pd.ExcelWriter(out, engine="openpyxl") as writer:
        df.to_excel(writer, index=False, sheet_name="Candidate Tracker")
        ws = writer.book["Candidate Tracker"]
        from openpyxl.styles import PatternFill, Font, Alignment
        header_fill = PatternFill("solid", fgColor="000000")
        ai_fill = PatternFill("solid", fgColor="C00000")
        white_font = Font(color="FFFFFF", bold=True)
        for cell in ws[1]:
            cell.fill = ai_fill if cell.column >= 33 else header_fill
            cell.font = white_font
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        for col in ws.columns:
            letter = col[0].column_letter
            max_len = max(len(str(cell.value or "")) for cell in col[:100])
            ws.column_dimensions[letter].width = min(max(max_len + 2, 12), 38)
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=True)
    out.seek(0)
    return StreamingResponse(
        out,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition":"attachment; filename=ShortlistAI_Candidate_Tracker.xlsx"}
    )
'''
    s = s.replace('\n@app.post("/export")\ndef export_excel', endpoint + '\n\n@app.post("/export")\ndef export_excel', 1)

main.write_text(s, encoding='utf-8')

# Frontend: add tracker download and duplicate-save feedback.
index = Path('static/index.html')
h = index.read_text(encoding='utf-8')
if 'id="exportTracker"' not in h:
    old = '''      <h3>Candidate Database</h3>\n      <div class="toolbar"><input id="searchCand" placeholder="Search candidate or skill"/><select id="filterStage"><option value="">All stages</option><option>Sourced</option><option>Screened</option><option>Interview</option><option>Offered</option><option>Joined</option><option>Rejected</option></select></div>'''
    new = '''      <div class="row" style="margin-bottom:12px"><h3 style="margin:0">Candidate Database</h3><button id="exportTracker" class="btn right">↓ Download Tracker</button></div>\n      <div class="toolbar"><input id="searchCand" placeholder="Search candidate or skill"/><select id="filterStage"><option value="">All stages</option><option>Sourced</option><option>Screened</option><option>Interview</option><option>Offered</option><option>Joined</option><option>Rejected</option></select></div>'''
    assert old in h
    h = h.replace(old, new, 1)

    anchor = '''$('searchCand').oninput=()=>loadCandidates();\n$('filterStage').onchange=()=>loadCandidates();\n'''
    handler = r'''$('searchCand').oninput=()=>loadCandidates();
$('filterStage').onchange=()=>loadCandidates();
$('exportTracker').onclick=async()=>{
  try{
    $('exportTracker').disabled=true;$('exportTracker').textContent='Preparing Excel…';
    const r=await fetch('/api/export/tracker');if(!r.ok)throw new Error('Could not export tracker');
    const blob=await r.blob(),a=document.createElement('a');a.href=URL.createObjectURL(blob);a.download='ShortlistAI_Candidate_Tracker.xlsx';a.click();URL.revokeObjectURL(a.href);toast('Tracker downloaded');
  }catch(e){toast(e.message)}finally{$('exportTracker').disabled=false;$('exportTracker').textContent='↓ Download Tracker'}
};
'''
    assert anchor in h
    h = h.replace(anchor, handler, 1)

# Make AI Match clearly report duplicates not saved.
old_status = "$('status').textContent=`Done — ${d.results.length} candidate(s) ranked.${d.errors?.length?' '+d.errors.length+' file(s) could not be read.':''}`;"
if old_status in h:
    new_status = "$('status').textContent=`Done — ${d.results.length} candidate(s) ranked.${d.duplicates_skipped?.length?' '+d.duplicates_skipped.length+' duplicate(s) not added again.':''}${d.errors?.length?' '+d.errors.length+' file(s) could not be read.':''}`;"
    h = h.replace(old_status, new_status, 1)

index.write_text(h, encoding='utf-8')
