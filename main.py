from __future__ import annotations
import io, re, sqlite3, os
from pathlib import Path
from typing import List, Dict, Any, Optional
from datetime import datetime

import pandas as pd
from fastapi import FastAPI, UploadFile, File, Form, HTTPException
from fastapi.responses import HTMLResponse, StreamingResponse
from fastapi.staticfiles import StaticFiles
from pydantic import BaseModel
from pypdf import PdfReader
from docx import Document
from sklearn.feature_extraction.text import TfidfVectorizer
from sklearn.metrics.pairwise import cosine_similarity

BASE_DIR = Path(__file__).resolve().parent
DB_PATH = os.getenv("SQLITE_PATH", str(BASE_DIR / "shortlistai.db"))
app = FastAPI(title="ShortlistAI ATS", version="1.0.0")
app.mount("/static", StaticFiles(directory=BASE_DIR / "static"), name="static")

SKILLS = [
    "python","fastapi","django","flask","java","spring boot","javascript","typescript","react","react.js","react 19","next.js","node.js","angular","vue",
    "html","css","tailwind","redux","rest api","graphql","microservices","sql","mysql","postgresql","mongodb","redis","elasticsearch","kafka","rabbitmq",
    "aws","azure","gcp","docker","kubernetes","terraform","jenkins","github actions","gitlab","ci/cd","linux","nginx",
    "generative ai","genai","agentic ai","llm","rag","langchain","langgraph","anthropic","claude","openai","mlflow","databricks","machine learning","deep learning",
    "pandas","numpy","scikit-learn","pytorch","tensorflow","computer vision","nlp","prompt engineering",
    "swift","ios","xcode","swiftui","uikit","android","kotlin","jetpack compose","flutter","react native",
    "selenium","playwright","cypress","appium","pytest","junit","testng","automation testing","api testing",
    "magento 2","adobe commerce","hcl commerce","shopify","salesforce","zoho crm","sap","oracle",
    "cybersecurity","google adk","oauth","rbac","sso","jwt","event streaming","erp integration",
    "product management","jira","figma","agile","scrum","stakeholder management","roadmap","analytics",
    "talent acquisition","it recruitment","technical recruitment","end-to-end recruitment","candidate sourcing","naukri","linkedin recruiter","screening","interview coordination","offer negotiation","stakeholder management","candidate engagement","recruitment mis"
]
ALIASES = {"react.js":"react","reactjs":"react","nodejs":"node.js","node":"node.js","gen ai":"genai","amazon web services":"aws","google cloud":"gcp","k8s":"kubernetes","postgres":"postgresql","restful api":"rest api","restful apis":"rest api"}
STOP = set("a an the and or of in on to for with from by as is are be been being this that these those you we our your will would should can may must role job candidate experience years year work working knowledge strong good excellent hands hand responsibilities requirements required preferred plus team teams ability skills skill using use used development developer engineer engineering design build building maintain maintaining".split())
STAGES = ["Sourced","Screened","Interview","Offered","Joined","Rejected"]

def db():
    con=sqlite3.connect(DB_PATH); con.row_factory=sqlite3.Row; return con

def init_db():
    con=db(); cur=con.cursor()
    cur.execute("""CREATE TABLE IF NOT EXISTS jobs(id INTEGER PRIMARY KEY AUTOINCREMENT,title TEXT NOT NULL,department TEXT,location TEXT,jd TEXT NOT NULL,status TEXT DEFAULT 'Open',created_at TEXT NOT NULL)""")
    cur.execute("""CREATE TABLE IF NOT EXISTS candidates(id INTEGER PRIMARY KEY AUTOINCREMENT,name TEXT NOT NULL,email TEXT,phone TEXT,experience REAL,skills TEXT,resume_text TEXT,source TEXT,notice_period TEXT,current_ctc TEXT,expected_ctc TEXT,job_id INTEGER,stage TEXT DEFAULT 'Sourced',ai_score REAL,rating TEXT,created_at TEXT NOT NULL,updated_at TEXT NOT NULL)""")
    cur.execute("""CREATE TABLE IF NOT EXISTS notes(id INTEGER PRIMARY KEY AUTOINCREMENT,candidate_id INTEGER NOT NULL,note TEXT NOT NULL,created_at TEXT NOT NULL)""")
    con.commit(); con.close()

@app.on_event("startup")
def startup(): init_db()

def normalize(text:str)->str:
    text=(text or "").lower().replace("–","-").replace("—","-")
    for src,dst in sorted(ALIASES.items(),key=lambda x:-len(x[0])):
        text=re.sub(r"(?<![a-z0-9])"+re.escape(src)+r"(?![a-z0-9])",dst,text)
    text=re.sub(r"[^a-z0-9+#./\- ]+"," ",text)
    return re.sub(r"\s+"," ",text).strip()

def extract_text(filename:str,data:bytes)->str:
    ext=Path(filename).suffix.lower()
    try:
        if ext==".pdf":
            reader=PdfReader(io.BytesIO(data)); return "\n".join((p.extract_text() or "") for p in reader.pages)
        if ext==".docx":
            doc=Document(io.BytesIO(data)); parts=[]
            parts += [p.text for p in doc.paragraphs if p.text and p.text.strip()]
            for table in doc.tables:
                for row in table.rows:
                    for cell in row.cells:
                        if cell.text.strip(): parts.append(cell.text.strip())
            for sec in doc.sections:
                parts += [p.text for p in sec.header.paragraphs if p.text and p.text.strip()]
                parts += [p.text for p in sec.footer.paragraphs if p.text and p.text.strip()]
            try:
                xml=[n.text for n in doc.element.body.iter() if n.tag.endswith('}t') and n.text]
                if len(" ".join(parts))<80: parts += xml
            except Exception: pass
            return "\n".join(parts)
        if ext in {".txt",".md"}: return data.decode("utf-8",errors="ignore")
    except Exception as e: raise ValueError(f"Could not read {filename}: {e}")
    raise ValueError(f"Unsupported format for {filename}. Use PDF, DOCX, TXT or MD.")

def find_skills(text:str)->List[str]:
    n=" "+normalize(text)+" "; out=[]
    for skill in SKILLS:
        s=normalize(skill)
        if re.search(r"(?<![a-z0-9])"+re.escape(s)+r"(?![a-z0-9])",n):
            c=ALIASES.get(s,s)
            if c not in out: out.append(c)
    return out

def parse_required_years(jd:str)->Optional[float]:
    n=normalize(jd)
    for p in [r"(\d+(?:\.\d+)?)\s*\+?\s*(?:to|-)\s*(\d+(?:\.\d+)?)\s*years?",r"(?:minimum|min|at least)\s*(\d+(?:\.\d+)?)\s*\+?\s*years?",r"(\d+(?:\.\d+)?)\s*\+\s*years?",r"(\d+(?:\.\d+)?)\s*years?\s*(?:of)?\s*experience"]:
        m=re.search(p,n)
        if m:return float(m.group(1))
    return None

def parse_candidate_years(resume:str)->Optional[float]:
    n=normalize(resume); vals=[]
    for p in [r"(\d+(?:\.\d+)?)\s*\+?\s*years?\s*(?:of)?\s*(?:overall|total|professional|relevant)?\s*experience",r"experience\s*(?:of|:)\s*(\d+(?:\.\d+)?)\s*\+?\s*years?"]:
        vals += [float(x) for x in re.findall(p,n)]
    vals=[v for v in vals if 0<=v<=50]; return max(vals) if vals else None

def detect_name(text:str,filename:str)->str:
    lines=[re.sub(r"\s+"," ",x).strip() for x in text.splitlines() if x.strip()]
    for line in lines[:12]:
        if 2<=len(line.split())<=5 and len(line)<60 and not any(k in line.lower() for k in ['resume','curriculum','summary','profile','email','phone','mobile']):
            if re.fullmatch(r"[A-Za-z .'-]+",line): return line.title()
    return Path(filename).stem.replace('_',' ').replace('-',' ').title()

def detect_email(text:str)->str:
    m=re.search(r"[A-Za-z0-9._%+-]+@[A-Za-z0-9.-]+\.[A-Za-z]{2,}",text); return m.group(0) if m else ""

def detect_phone(text:str)->str:
    m=re.search(r"(?:\+?91[\s-]?)?[6-9]\d{9}",re.sub(r"[() ]","",text)); return m.group(0) if m else ""

def semantic_scores(jd:str,resumes:List[str])->List[float]:
    try:
        mat=TfidfVectorizer(stop_words="english",ngram_range=(1,2),max_features=12000).fit_transform([normalize(jd)]+[normalize(r) for r in resumes])
        return [float(x) for x in cosine_similarity(mat[0:1],mat[1:]).flatten()]
    except Exception:return [0.0]*len(resumes)

def rating(score:float)->str:
    return "Strong" if score>=75 else "Average" if score>=55 else "Weak"

def score_resume(jd:str,resume:str,sim:float)->Dict[str,Any]:
    req=find_skills(jd); got=find_skills(resume); matched=[s for s in req if s in got]; missing=[s for s in req if s not in got]
    cov=len(matched)/len(req) if req else .5; reqy=parse_required_years(jd); cy=parse_candidate_years(resume)
    exp=.75 if reqy is None and cy is not None else .55 if reqy is None else .45 if cy is None else min(cy/max(reqy,1),1)
    n=normalize(resume); tail=n[int(len(n)*.55):]; recent=sum(1 for s in matched if normalize(s) in tail)/len(matched) if matched else 0
    total=round(max(0,min(100,45*sim+35*cov+10*exp+10*recent)),1)
    risks=[]
    if missing: risks.append("Missing/unclear JD skills: "+", ".join(missing[:6]))
    if cy is None: risks.append("Total experience could not be reliably extracted; verify manually.")
    if reqy and cy is not None and cy<reqy: risks.append(f"Resume states ~{cy:g} years vs JD minimum ~{reqy:g} years.")
    if not risks: risks.append("No major resume-level gap detected; validate depth during screening.")
    questions=[f"Describe your most recent hands-on work with {s}. What did you personally own?" for s in (missing[:2]+matched[:2])][:4] or ["Walk me through your most relevant recent project and what you personally owned."]
    return {"score":total,"rating":rating(total),"semantic_similarity":round(sim*100,1),"skill_coverage":round(cov*100,1),"experience_fit":round(exp*100,1),"recent_evidence":round(recent*100,1),"candidate_years":cy,"required_years":reqy,"matched_skills":matched,"missing_skills":missing,"recent_skills":[s for s in matched if normalize(s) in tail],"risks":risks,"evidence":[f"Matched {len(matched)}/{len(req) or len(matched)} detected JD skills."],"screening_questions":questions}

@app.get("/",response_class=HTMLResponse)
def home(): return (BASE_DIR/"static"/"index.html").read_text(encoding="utf-8")

class JobIn(BaseModel):
    title:str; department:str=""; location:str=""; jd:str; status:str="Open"
class CandidateIn(BaseModel):
    name:str; email:str=""; phone:str=""; experience:Optional[float]=None; skills:str=""; source:str=""; notice_period:str=""; current_ctc:str=""; expected_ctc:str=""; job_id:Optional[int]=None; stage:str="Sourced"
class StageIn(BaseModel): stage:str
class NoteIn(BaseModel): note:str
class ExportPayload(BaseModel): results:List[Dict[str,Any]]

@app.get("/api/stats")
def stats():
    con=db(); c=con.cursor(); out={}
    out['jobs']=c.execute("SELECT COUNT(*) FROM jobs WHERE status='Open'").fetchone()[0]
    out['candidates']=c.execute("SELECT COUNT(*) FROM candidates").fetchone()[0]
    out['interviews']=c.execute("SELECT COUNT(*) FROM candidates WHERE stage='Interview'").fetchone()[0]
    out['offers']=c.execute("SELECT COUNT(*) FROM candidates WHERE stage='Offered'").fetchone()[0]
    out['joined']=c.execute("SELECT COUNT(*) FROM candidates WHERE stage='Joined'").fetchone()[0]
    con.close(); return out

@app.get("/api/jobs")
def list_jobs():
    con=db(); rows=[dict(x) for x in con.execute("SELECT * FROM jobs ORDER BY id DESC")]; con.close(); return rows
@app.post("/api/jobs")
def create_job(x:JobIn):
    if len(x.jd.strip())<20: raise HTTPException(400,"Please enter a fuller JD")
    con=db(); cur=con.cursor(); cur.execute("INSERT INTO jobs(title,department,location,jd,status,created_at) VALUES(?,?,?,?,?,?)",(x.title,x.department,x.location,x.jd,x.status,datetime.utcnow().isoformat())); con.commit(); i=cur.lastrowid; con.close(); return {"id":i}
@app.delete("/api/jobs/{job_id}")
def delete_job(job_id:int):
    con=db(); con.execute("DELETE FROM jobs WHERE id=?",(job_id,)); con.commit(); con.close(); return {"ok":True}

@app.get("/api/candidates")
def list_candidates(job_id:Optional[int]=None,stage:Optional[str]=None,q:Optional[str]=None):
    sql="SELECT c.*,j.title job_title FROM candidates c LEFT JOIN jobs j ON j.id=c.job_id WHERE 1=1"; args=[]
    if job_id: sql+=" AND c.job_id=?"; args.append(job_id)
    if stage: sql+=" AND c.stage=?"; args.append(stage)
    if q: sql+=" AND (c.name LIKE ? OR c.email LIKE ? OR c.skills LIKE ?)"; v=f"%{q}%"; args += [v,v,v]
    sql+=" ORDER BY COALESCE(c.ai_score,0) DESC,c.id DESC"
    con=db(); rows=[dict(x) for x in con.execute(sql,args)]; con.close(); return rows
@app.post("/api/candidates")
def create_candidate(x:CandidateIn):
    if x.stage not in STAGES: raise HTTPException(400,"Invalid stage")
    now=datetime.utcnow().isoformat(); con=db(); cur=con.cursor(); cur.execute("""INSERT INTO candidates(name,email,phone,experience,skills,source,notice_period,current_ctc,expected_ctc,job_id,stage,created_at,updated_at) VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?)""",(x.name,x.email,x.phone,x.experience,x.skills,x.source,x.notice_period,x.current_ctc,x.expected_ctc,x.job_id,x.stage,now,now)); con.commit(); i=cur.lastrowid; con.close(); return {"id":i}
@app.patch("/api/candidates/{candidate_id}/stage")
def update_stage(candidate_id:int,x:StageIn):
    if x.stage not in STAGES: raise HTTPException(400,"Invalid stage")
    con=db(); con.execute("UPDATE candidates SET stage=?,updated_at=? WHERE id=?",(x.stage,datetime.utcnow().isoformat(),candidate_id)); con.commit(); con.close(); return {"ok":True}
@app.delete("/api/candidates/{candidate_id}")
def delete_candidate(candidate_id:int):
    con=db(); con.execute("DELETE FROM notes WHERE candidate_id=?",(candidate_id,)); con.execute("DELETE FROM candidates WHERE id=?",(candidate_id,)); con.commit(); con.close(); return {"ok":True}
@app.get("/api/candidates/{candidate_id}/notes")
def get_notes(candidate_id:int):
    con=db(); rows=[dict(x) for x in con.execute("SELECT * FROM notes WHERE candidate_id=? ORDER BY id DESC",(candidate_id,))]; con.close(); return rows
@app.post("/api/candidates/{candidate_id}/notes")
def add_note(candidate_id:int,x:NoteIn):
    if not x.note.strip(): raise HTTPException(400,"Note is empty")
    con=db(); con.execute("INSERT INTO notes(candidate_id,note,created_at) VALUES(?,?,?)",(candidate_id,x.note.strip(),datetime.utcnow().isoformat())); con.commit(); con.close(); return {"ok":True}

@app.post("/analyze")
async def analyze(jd:str=Form(...),resumes:List[UploadFile]=File(...),job_id:Optional[int]=Form(None),save_to_ats:bool=Form(False)):
    if len(jd.strip())<50: raise HTTPException(400,"Please paste a more complete job description.")
    if not resumes: raise HTTPException(400,"Upload at least one resume.")
    texts=[]; names=[]; errors=[]
    for f in resumes[:50]:
        data=await f.read()
        try:
            txt=extract_text(f.filename or "resume",data)
            if len(txt.strip())<80: raise ValueError("Very little readable text was extracted")
            texts.append(txt); names.append(f.filename or "resume")
        except Exception as e: errors.append(str(e))
    if not texts: raise HTTPException(400,"No resumes could be read. "+"; ".join(errors))
    sims=semantic_scores(jd,texts); results=[]; now=datetime.utcnow().isoformat()
    con=db() if save_to_ats else None
    for txt,name,sim in zip(texts,names,sims):
        s=score_resume(jd,txt,sim); s.update({"candidate":detect_name(txt,name),"file":name})
        results.append(s)
        if con:
            skills=", ".join(find_skills(txt))
            con.execute("""INSERT INTO candidates(name,email,phone,experience,skills,resume_text,source,job_id,stage,ai_score,rating,created_at,updated_at) VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?)""",(s['candidate'],detect_email(txt),detect_phone(txt),s['candidate_years'],skills,txt,'Resume upload',job_id,'Sourced',s['score'],s['rating'],now,now))
    if con: con.commit(); con.close()
    results.sort(key=lambda x:x['score'],reverse=True); req=find_skills(jd)
    return {"summary":{"candidates":len(results),"required_skills":req,"required_years":parse_required_years(jd),"strong":sum(r['rating']=='Strong' for r in results),"average":sum(r['rating']=='Average' for r in results),"weak":sum(r['rating']=='Weak' for r in results)},"results":results,"errors":errors,"methodology":"45% semantic similarity + 35% JD skill coverage + 10% experience fit + 10% recent skill evidence. Use as recruiter decision support, not an autonomous hiring decision."}

@app.post("/export")
def export_excel(payload:ExportPayload):
    rows=[]
    for i,r in enumerate(payload.results,1): rows.append({"Rank":i,"Candidate":r.get('candidate'),"Score":r.get('score'),"Rating":r.get('rating'),"Candidate Years":r.get('candidate_years'),"Matched Skills":", ".join(r.get('matched_skills',[])),"Missing Skills":", ".join(r.get('missing_skills',[]))})
    out=io.BytesIO(); pd.DataFrame(rows).to_excel(out,index=False); out.seek(0)
    return StreamingResponse(out,media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",headers={"Content-Disposition":"attachment; filename=ShortlistAI_results.xlsx"})
